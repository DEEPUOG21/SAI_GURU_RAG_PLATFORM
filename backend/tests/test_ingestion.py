import io
import zipfile

import pytest
from docx import Document

from app.ingestion.parsing import Section, chunk_sections, parse, validate_file


@pytest.mark.parametrize(
    "name,content,expected",
    [
        ("policy.md", b"# Leave\n\nStaff have 24 days.", "24 days"),
        ("data.csv", b"name,days\nLeave,24", "days: 24"),
        ("data.json", b'[{"days":24}]', '"days": 24'),
        ("policy.html", b"<h2>Leave</h2><p>24 days</p><script>bad()</script>", "24 days"),
        ("policy.txt", b"Staff have 24 days.", "24 days"),
    ],
)
def test_parsers(name, content, expected):
    validate_file(name, content)
    sections = parse(name, content)
    assert expected in sections[0].text
    assert "bad()" not in sections[0].text


def test_docx():
    doc = Document()
    doc.add_heading("Expenses", level=1)
    doc.add_paragraph("Submit within 30 days.")
    buffer = io.BytesIO()
    doc.save(buffer)
    validate_file("test.docx", buffer.getvalue())
    assert parse("test.docx", buffer.getvalue())[0].heading == "Expenses"


def test_structure_and_long_sentences():
    chunks = chunk_sections(
        [Section("One sentence. Second sentence.", "Leave", 3), Section("word " * 500, "Security", 4)],
        max_words=50,
    )
    assert chunks[0].heading == "Leave" and chunks[0].page == 3
    assert all(len(c.text.split()) <= 50 for c in chunks)
    assert sum(len(c.text.split()) for c in chunks) == 504


@pytest.mark.parametrize(
    "name,content",
    [("../x.md", b"x"), ("x.exe", b"x"), ("x.pdf", b"fake"), ("x.txt", b""), ("x.docx", b"PKfake")],
)
def test_reject_invalid_uploads(name, content):
    with pytest.raises(ValueError):
        validate_file(name, content)


def test_reject_fake_docx():
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("file.txt", "fake")
    with pytest.raises(ValueError):
        validate_file("x.docx", stream.getvalue())


def test_pdf_preserves_real_page_metadata():
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 700 Td (Leave is 24 days.) Tj ET")
    page[NameObject("/Contents")] = stream
    buffer = io.BytesIO()
    writer.write(buffer)
    validate_file("policy.pdf", buffer.getvalue())
    sections = parse("policy.pdf", buffer.getvalue())
    assert sections[0].page == 1 and "24 days" in sections[0].text
